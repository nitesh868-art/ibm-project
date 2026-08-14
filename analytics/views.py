"""Analytics Views"""
from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from datetime import timedelta


@login_required
def analytics_home(request):
    return render(request, 'analytics/home.html')


@login_required
def study_analytics(request):
    from studyplanner.models import DailyStudyLog, StudySubject
    today = timezone.now().date()
    week_ago = today - timedelta(days=7)
    month_ago = today - timedelta(days=30)

    # Weekly study hours
    weekly_logs = DailyStudyLog.objects.filter(student=request.user, date__gte=week_ago)
    daily_hours = []
    for i in range(7):
        day = today - timedelta(days=6-i)
        day_logs = weekly_logs.filter(date=day)
        total = sum(l.duration_minutes for l in day_logs)
        daily_hours.append({'day': day.strftime('%a %d'), 'hours': round(total/60, 1)})

    # Subject-wise time
    subjects = StudySubject.objects.filter(student=request.user)
    subject_data = []
    for s in subjects:
        logs = DailyStudyLog.objects.filter(student=request.user, subject=s)
        total_mins = sum(l.duration_minutes for l in logs)
        subject_data.append({'name': s.name, 'hours': round(total_mins/60, 1), 'color': s.color})

    return render(request, 'analytics/study.html', {
        'daily_hours': daily_hours,
        'subject_data': subject_data,
    })


@login_required
def placement_analytics(request):
    from quiz.models import TestAttempt
    attempts = TestAttempt.objects.filter(student=request.user, is_completed=True)
    scores = [{'test': a.test.title[:30], 'score': float(a.percentage)} for a in attempts[:10]]
    return render(request, 'analytics/placement.html', {'scores': scores})


@login_required
def export_pdf(request):
    """Export analytics report as PDF."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        import io

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        story = [
            Paragraph(f'Analytics Report — {request.user.get_full_name()}', styles['Title']),
            Spacer(1, 20),
            Paragraph(f'Generated on: {timezone.now().strftime("%d %B %Y")}', styles['Normal']),
        ]
        doc.build(story)
        buffer.seek(0)
        response = HttpResponse(buffer.read(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="analytics_report.pdf"'
        return response
    except Exception as e:
        from django.contrib import messages
        messages.error(request, f'Export failed: {e}')
        return render(request, 'analytics/home.html')


@login_required
def export_excel(request):
    """Export analytics data as Excel."""
    try:
        import openpyxl
        from openpyxl import Workbook
        import io

        wb = Workbook()
        ws = wb.active
        ws.title = 'Analytics'
        ws['A1'] = 'AI Placement Portal — Analytics Report'
        ws['A2'] = f'Student: {request.user.get_full_name()}'
        ws['A3'] = f'Date: {timezone.now().strftime("%d %B %Y")}'

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="analytics.xlsx"'
        return response
    except Exception as e:
        from django.contrib import messages
        messages.error(request, f'Export failed: {e}')
        return render(request, 'analytics/home.html')


@login_required
def chart_data(request):
    """API endpoint for chart data."""
    from studyplanner.models import DailyStudyLog
    from quiz.models import TestAttempt
    today = timezone.now().date()

    # Last 7 days study data
    weekly_data = []
    for i in range(7):
        day = today - timedelta(days=6-i)
        logs = DailyStudyLog.objects.filter(student=request.user, date=day)
        total = sum(l.duration_minutes for l in logs)
        weekly_data.append(round(total/60, 1))

    # Last 10 test scores
    scores = list(
        TestAttempt.objects.filter(student=request.user, is_completed=True)
        .order_by('-started_at')[:10]
        .values_list('percentage', flat=True)
    )

    return JsonResponse({
        'weekly_hours': weekly_data,
        'test_scores': [float(s) for s in scores],
        'days': [(today - timedelta(days=6-i)).strftime('%a') for i in range(7)],
    })
